from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook

from compatibilizabim.sinapi import SinapiDatabase, SinapiImporter, SinapiService


def _zip_wide(tmp_path: Path) -> Path:
    xlsx = tmp_path / "SINAPI_202607.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Composicoes"
    ws.append(["Código da Composição", "Descrição", "Unidade", "SP", "RJ"])
    ws.append([103328, "ALVENARIA DE VEDACAO DE BLOCOS CERAMICOS 14CM", "M2", 82.43, 86.10])
    ws.append([94965, "CONCRETO FCK 25 MPA", "M3", 612.55, 635.20])
    ins = wb.create_sheet("Insumos")
    ins.append(["Código do Insumo", "Descrição do Insumo", "Unidade", "SP", "RJ"])
    ins.append([7258, "TIJOLO CERAMICO", "UN", 1.25, 1.31])
    wb.save(xlsx)
    archive = tmp_path / "sinapi.zip"
    with ZipFile(archive, "w") as zf:
        zf.write(xlsx, xlsx.name)
    return archive


def _zip_long(tmp_path: Path) -> Path:
    xlsx = tmp_path / "relatorio.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Custos"
    ws.append(["Tipo", "Código", "Descrição", "Unidade", "UF", "Custo"])
    ws.append(["Composição", "10001", "REVESTIMENTO CERAMICO PAREDE", "M2", "SP", "55,32"])
    ws.append(["Insumo", "20001", "ARGAMASSA", "KG", "SP", "2,45"])
    wb.save(xlsx)
    archive = tmp_path / "long.zip"
    with ZipFile(archive, "w") as zf:
        zf.write(xlsx, xlsx.name)
    return archive


def test_import_wide_zip_versions_and_searches_by_uf(tmp_path: Path):
    db = SinapiDatabase(tmp_path / "sinapi.sqlite3")
    service = SinapiService(db)
    archive = _zip_wide(tmp_path)

    release = SinapiImporter(db).import_archive(archive, competence="2026-07")
    again = SinapiImporter(db).import_archive(archive, competence="2026-07")

    assert release.release_id == again.release_id
    assert release.competence == "2026-07"
    assert release.source_sha256 == again.source_sha256
    assert db.count_items(release.release_id, kind="composition") == 4
    assert db.count_items(release.release_id, kind="input") == 2

    hits = service.search_compositions(release.release_id, uf="SP", query="alvenaria ceramicos", limit=5)
    assert hits[0]["code"] == "103328"
    assert hits[0]["unit"] == "m2"
    assert hits[0]["price"] == 82.43
    assert hits[0]["uf"] == "SP"


def test_import_long_format_detects_kind_uf_and_decimal_comma(tmp_path: Path):
    db = SinapiDatabase(tmp_path / "sinapi.sqlite3")
    release = SinapiImporter(db).import_archive(_zip_long(tmp_path), competence="2026-06")

    comp = db.get_item(release.release_id, kind="composition", uf="SP", code="10001")
    inp = db.get_item(release.release_id, kind="input", uf="SP", code="20001")
    assert comp is not None and comp.price == 55.32 and comp.unit == "m2"
    assert inp is not None and inp.price == 2.45 and inp.unit == "kg"


def test_build_pricebook_from_confirmed_sinapi_mapping(tmp_path: Path):
    db = SinapiDatabase(tmp_path / "sinapi.sqlite3")
    release = SinapiImporter(db).import_archive(_zip_wide(tmp_path), competence="2026-07")
    service = SinapiService(db)

    book = service.build_pricebook(
        release.release_id,
        uf="SP",
        bdi_percent=22.0,
        mappings=[{
            "composition_code": "103328",
            "ifc_class": "IfcWall",
            "quantity_kind": "area",
            "quantity_name_contains": "NetSideArea",
        }],
    )

    assert book.source == "SINAPI"
    assert book.reference_date == "2026-07/SP"
    assert book.compositions[0].code == "SINAPI-103328"
    assert book.compositions[0].base_unit_cost == 82.43
    assert book.mappings[0].composition_code == "SINAPI-103328"


def test_compare_releases_reports_price_variation(tmp_path: Path):
    db = SinapiDatabase(tmp_path / "sinapi.sqlite3")
    importer = SinapiImporter(db)
    old_zip = _zip_wide(tmp_path)
    old = importer.import_archive(old_zip, competence="2026-06")

    # Same codes, changed SP price in another physical source file.
    xlsx = tmp_path / "SINAPI_202607_B.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "Composicoes"
    ws.append(["Código", "Descrição", "Unidade", "SP"])
    ws.append([103328, "ALVENARIA DE VEDACAO DE BLOCOS CERAMICOS 14CM", "M2", 84.08])
    wb.save(xlsx)
    new_zip = tmp_path / "new.zip"
    with ZipFile(new_zip, "w") as zf: zf.write(xlsx, xlsx.name)
    new = importer.import_archive(new_zip, competence="2026-07")

    report = SinapiService(db).compare_releases(old.release_id, new.release_id, uf="SP", codes=["103328"])
    assert report["rows"][0]["old_price"] == 82.43
    assert report["rows"][0]["new_price"] == 84.08
    assert round(report["rows"][0]["variation_percent"], 2) == 2.00

def test_invalid_import_does_not_leave_empty_release(tmp_path: Path):
    bad = tmp_path / 'bad.zip'
    with ZipFile(bad, 'w') as zf:
        zf.writestr('readme.txt', 'sem xlsx')
    db = SinapiDatabase(tmp_path / 'sinapi.sqlite3')
    try:
        SinapiImporter(db).import_archive(bad, competence='2026-07')
    except ValueError:
        pass
    else:
        raise AssertionError('import deveria falhar')
    assert db.list_releases() == []

def test_import_infers_uf_from_sheet_when_price_is_single_column(tmp_path: Path):
    xlsx = tmp_path / 'custos.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'SP'
    ws.append(['Código da Composição','Descrição','Unidade','Preço'])
    ws.append([777,'PAREDE TESTE','M2',91.5]); wb.save(xlsx)
    db = SinapiDatabase(tmp_path/'db.sqlite3')
    release = SinapiImporter(db).import_archive(xlsx, competence='2026-08')
    item = db.get_item(release.release_id, kind='composition', uf='SP', code='777')
    assert item is not None and item.price == 91.5
