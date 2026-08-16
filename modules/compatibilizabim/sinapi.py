from __future__ import annotations

import hashlib
import io
import math
import re
import sqlite3
import tempfile
import unicodedata
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence
from urllib.parse import urlparse
from urllib.request import Request, urlopen
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .budget import Composition, CostComponent, CostMapping, PriceBook

UFS = {
    "AC","AL","AP","AM","BA","CE","DF","ES","GO","MA","MT","MS","MG","PA","PB","PR","PE","PI","RJ","RN","RS","RO","RR","SC","SP","SE","TO"
}

CAIXA_XLSX_URL_TEMPLATE = (
    "https://www.caixa.gov.br/Downloads/sinapi-proposta-formato-divulgacao-mensal/"
    "SINAPI-{year}-{month}-formato-xlsx.zip"
)


@dataclass(frozen=True, slots=True)
class SinapiOfficialReference:
    competence: str
    published_on: str
    provider: str
    verification_url: str
    landing_url: str


LATEST_VERIFIED_SINAPI = SinapiOfficialReference(
    competence="2026-07",
    published_on="2026-08-11",
    provider="CAIXA/IBGE",
    verification_url="https://www.ibge.gov.br/calendario/conjunturais.html",
    landing_url="https://www.caixa.gov.br/poder-publico/modernizacao-gestao/sinapi/Paginas/default.aspx",
)


def official_xlsx_url(competence: str) -> str:
    if not re.fullmatch(r"\d{4}-\d{2}", competence):
        raise ValueError("Competência deve estar no formato YYYY-MM")
    year, month = competence.split("-")
    return CAIXA_XLSX_URL_TEMPLATE.format(year=year, month=month, competence=competence)


@dataclass(frozen=True, slots=True)
class DownloadedSinapiArchive:
    path: Path
    competence: str
    source_url: str
    sha256: str
    fetched_at: str
    xlsx_members: tuple[str, ...]


class OfficialSinapiDownloader:
    """Downloads a CAIXA monthly XLSX archive and validates it before publication.

    The default configuration trusts only the official CAIXA host. Tests may inject a
    loopback URL/template explicitly; no third-party fallback is performed.
    """

    def __init__(
        self,
        *,
        base_url_template: str = CAIXA_XLSX_URL_TEMPLATE,
        trusted_hosts: set[str] | None = None,
        timeout: float = 90.0,
        max_bytes: int = 1024 * 1024 * 1024,
    ) -> None:
        self.base_url_template = base_url_template
        self.trusted_hosts = trusted_hosts or {"www.caixa.gov.br", "caixa.gov.br"}
        self.timeout = float(timeout)
        self.max_bytes = int(max_bytes)

    def url_for(self, competence: str) -> str:
        if not re.fullmatch(r"\d{4}-\d{2}", competence):
            raise ValueError("Competência deve estar no formato YYYY-MM")
        year, month = competence.split("-")
        url = self.base_url_template.format(year=year, month=month, competence=competence)
        parsed = urlparse(url)
        if parsed.hostname not in self.trusted_hosts:
            raise ValueError("Download SINAPI recusado: host não oficial/não confiável")
        if parsed.scheme not in {"https", "http"}:
            raise ValueError("Download SINAPI recusado: esquema de URL inválido")
        if parsed.scheme != "https" and parsed.hostname not in {"127.0.0.1", "localhost", "::1"}:
            raise ValueError("Download SINAPI oficial exige HTTPS")
        return url

    def download(self, competence: str, destination_dir: Path) -> DownloadedSinapiArchive:
        url = self.url_for(competence)
        destination_dir = Path(destination_dir)
        destination_dir.mkdir(parents=True, exist_ok=True)
        final = destination_dir / f"SINAPI-{competence}-formato-xlsx.zip"
        temp = destination_dir / f".{final.name}.part"
        temp.unlink(missing_ok=True)
        sha = hashlib.sha256()
        total = 0
        try:
            request = Request(url, headers={"User-Agent": "CompatibilizaBIM/1.3 SINAPI-official-updater"})
            with urlopen(request, timeout=self.timeout) as response, temp.open("wb") as handle:
                while True:
                    chunk = response.read(1024 * 1024)
                    if not chunk:
                        break
                    total += len(chunk)
                    if total > self.max_bytes:
                        raise ValueError("Arquivo SINAPI excede o limite de segurança")
                    sha.update(chunk)
                    handle.write(chunk)
            members = self._validate_archive(temp)
            temp.replace(final)
            return DownloadedSinapiArchive(
                path=final, competence=competence, source_url=url, sha256=sha.hexdigest(),
                fetched_at=datetime.now(timezone.utc).isoformat(), xlsx_members=members,
            )
        except Exception:
            temp.unlink(missing_ok=True)
            raise

    def _validate_archive(self, path: Path) -> tuple[str, ...]:
        try:
            with ZipFile(path) as zf:
                members: list[str] = []
                expanded = 0
                for info in zf.infolist():
                    name = info.filename.replace("\\", "/")
                    parts = Path(name).parts
                    if name.startswith("/") or ".." in parts:
                        raise ValueError("ZIP SINAPI contém caminho inseguro")
                    if info.flag_bits & 0x1:
                        raise ValueError("ZIP SINAPI criptografado não é suportado")
                    expanded += int(info.file_size)
                    if expanded > self.max_bytes * 4:
                        raise ValueError("ZIP SINAPI excede o limite descompactado")
                    if not info.is_dir() and name.lower().endswith(".xlsx"):
                        members.append(name)
                if not members:
                    raise ValueError("ZIP SINAPI inválido: nenhum XLSX encontrado")
                return tuple(sorted(members))
        except BadZipFile as exc:
            raise ValueError("ZIP SINAPI inválido") from exc


def _norm(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-zA-Z0-9]+", " ", text).strip().lower()
    return re.sub(r"\s+", " ", text)


def _code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().replace(".0", "")


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return float(value)
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if not text:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _uf_hint(*values: Any) -> str | None:
    for value in values:
        text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().upper()
        for token in re.findall(r"\b[A-Z]{2}\b", text):
            if token in UFS:
                return token
    return None


def normalize_unit(unit: Any) -> str:
    key = _norm(unit).replace(" ", "")
    mapping = {
        "m": "m", "metro": "m", "metros": "m",
        "m2": "m2", "m²": "m2", "metroquadrado": "m2",
        "m3": "m3", "m³": "m3", "metrocubico": "m3",
        "un": "un", "und": "un", "unidade": "un", "unidades": "un",
        "kg": "kg", "quilograma": "kg", "h": "h", "hora": "h",
        "l": "L", "litro": "L", "t": "t", "ton": "t",
    }
    return mapping.get(key, str(unit or "").strip().lower())


@dataclass(frozen=True, slots=True)
class SinapiRelease:
    release_id: str
    competence: str
    source_filename: str
    source_sha256: str
    imported_at: str
    source_url: str | None = None
    source_provider: str | None = None
    published_on: str | None = None
    fetched_at: str | None = None


@dataclass(frozen=True, slots=True)
class SinapiItem:
    release_id: str
    kind: str
    uf: str
    code: str
    description: str
    unit: str
    price: float
    workbook: str
    sheet: str


class SinapiDatabase:
    SCHEMA_VERSION = 2

    def __init__(self, path: Path) -> None:
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._init()

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.path)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA journal_mode=WAL")
        return conn

    def _init(self) -> None:
        with self._connect() as conn:
            conn.executescript("""
            CREATE TABLE IF NOT EXISTS metadata(key TEXT PRIMARY KEY, value TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS releases(
              release_id TEXT PRIMARY KEY,
              competence TEXT NOT NULL,
              source_filename TEXT NOT NULL,
              source_sha256 TEXT NOT NULL,
              imported_at TEXT NOT NULL,
              source_url TEXT,
              source_provider TEXT,
              published_on TEXT,
              fetched_at TEXT,
              UNIQUE(competence, source_sha256)
            );
            CREATE TABLE IF NOT EXISTS items(
              release_id TEXT NOT NULL,
              kind TEXT NOT NULL,
              uf TEXT NOT NULL,
              code TEXT NOT NULL,
              description TEXT NOT NULL,
              unit TEXT NOT NULL,
              price REAL NOT NULL,
              workbook TEXT NOT NULL,
              sheet TEXT NOT NULL,
              PRIMARY KEY(release_id, kind, uf, code),
              FOREIGN KEY(release_id) REFERENCES releases(release_id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS idx_items_search ON items(release_id, kind, uf, code);
            """)
            columns = {row[1] for row in conn.execute("PRAGMA table_info(releases)").fetchall()}
            for name, ddl in (
                ("source_url", "TEXT"),
                ("source_provider", "TEXT"),
                ("published_on", "TEXT"),
                ("fetched_at", "TEXT"),
            ):
                if name not in columns:
                    conn.execute(f"ALTER TABLE releases ADD COLUMN {name} {ddl}")
            conn.execute("INSERT OR REPLACE INTO metadata(key,value) VALUES('schema_version',?)", (str(self.SCHEMA_VERSION),))

    def get_release_by_source(self, competence: str, sha256: str) -> SinapiRelease | None:
        with self._connect() as conn:
            row = conn.execute("SELECT * FROM releases WHERE competence=? AND source_sha256=?", (competence, sha256)).fetchone()
        return SinapiRelease(**dict(row)) if row else None

    def insert_release(self, release: SinapiRelease) -> None:
        with self._connect() as conn:
            conn.execute(
                "INSERT INTO releases(release_id,competence,source_filename,source_sha256,imported_at,source_url,source_provider,published_on,fetched_at) VALUES(?,?,?,?,?,?,?,?,?)",
                (release.release_id, release.competence, release.source_filename, release.source_sha256, release.imported_at,
                 release.source_url, release.source_provider, release.published_on, release.fetched_at),
            )

    def insert_items(self, items: Iterable[SinapiItem]) -> None:
        rows = [tuple(asdict(item).values()) for item in items]
        if not rows:
            return
        with self._connect() as conn:
            conn.executemany(
                "INSERT OR REPLACE INTO items(release_id,kind,uf,code,description,unit,price,workbook,sheet) VALUES(?,?,?,?,?,?,?,?,?)",
                rows,
            )

    def list_releases(self) -> list[SinapiRelease]:
        with self._connect() as conn:
            rows = conn.execute("SELECT * FROM releases ORDER BY competence DESC, imported_at DESC").fetchall()
        return [SinapiRelease(**dict(row)) for row in rows]

    def count_items(self, release_id: str, *, kind: str | None = None) -> int:
        sql = "SELECT COUNT(*) FROM items WHERE release_id=?"; args: list[Any] = [release_id]
        if kind:
            sql += " AND kind=?"; args.append(kind)
        with self._connect() as conn:
            return int(conn.execute(sql, args).fetchone()[0])

    def get_item(self, release_id: str, *, kind: str, uf: str, code: str) -> SinapiItem | None:
        with self._connect() as conn:
            row = conn.execute("SELECT * FROM items WHERE release_id=? AND kind=? AND uf=? AND code=?", (release_id, kind, uf.upper(), str(code))).fetchone()
        return SinapiItem(**dict(row)) if row else None

    def items(self, release_id: str, *, kind: str, uf: str) -> list[SinapiItem]:
        with self._connect() as conn:
            rows = conn.execute("SELECT * FROM items WHERE release_id=? AND kind=? AND uf=? ORDER BY code", (release_id, kind, uf.upper())).fetchall()
        return [SinapiItem(**dict(row)) for row in rows]


class SinapiImporter:
    def __init__(self, db: SinapiDatabase) -> None:
        self.db = db

    def import_archive(
        self, path: Path, *, competence: str, source_url: str | None = None,
        source_provider: str | None = None, published_on: str | None = None,
        fetched_at: str | None = None,
    ) -> SinapiRelease:
        path = Path(path)
        if not re.fullmatch(r"\d{4}-\d{2}", competence):
            raise ValueError("Competência deve estar no formato YYYY-MM")
        if not path.is_file():
            raise FileNotFoundError(path)
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        existing = self.db.get_release_by_source(competence, sha)
        if existing:
            return existing
        release = SinapiRelease(
            release_id=f"SINAPI-{competence}-{sha[:12]}", competence=competence,
            source_filename=path.name, source_sha256=sha,
            imported_at=datetime.now(timezone.utc).isoformat(), source_url=source_url,
            source_provider=source_provider, published_on=published_on, fetched_at=fetched_at,
        )
        items: list[SinapiItem] = []
        for workbook_path, display_name in self._workbooks(path):
            items.extend(self._parse_workbook(workbook_path, display_name, release.release_id))
        if not items:
            raise ValueError("Nenhum relatório de insumos/composições reconhecido no arquivo")
        self.db.insert_release(release)
        self.db.insert_items(items)
        return release

    def _workbooks(self, path: Path):
        if path.suffix.lower() == ".xlsx":
            yield path, path.name
            return
        if path.suffix.lower() != ".zip":
            raise ValueError("Use um arquivo .zip ou .xlsx")
        with tempfile.TemporaryDirectory(prefix="sinapi-") as tmp:
            with ZipFile(path) as zf:
                for info in zf.infolist():
                    if info.is_dir() or not info.filename.lower().endswith(".xlsx"):
                        continue
                    name = Path(info.filename).name
                    out = Path(tmp) / f"{len(list(Path(tmp).glob('*.xlsx'))):04d}-{name}"
                    out.write_bytes(zf.read(info))
                    yield out, name

    def _parse_workbook(self, path: Path, display_name: str, release_id: str) -> list[SinapiItem]:
        wb = load_workbook(path, read_only=True, data_only=True)
        result: list[SinapiItem] = []
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            sample = []
            for _ in range(40):
                try: sample.append(next(rows))
                except StopIteration: break
            header_idx, headers = self._find_header(sample)
            if headers is None:
                continue
            kind_hint = self._kind_hint(ws.title, display_name, headers)
            uf_hint = _uf_hint(ws.title, display_name)
            for row in sample[header_idx + 1:]:
                result.extend(self._parse_row(row, headers, kind_hint, release_id, display_name, ws.title, uf_hint))
            for row in rows:
                result.extend(self._parse_row(row, headers, kind_hint, release_id, display_name, ws.title, uf_hint))
        wb.close()
        return result

    def _find_header(self, rows: Sequence[Sequence[Any]]) -> tuple[int, list[str] | None]:
        for idx, row in enumerate(rows):
            headers = [_norm(v) for v in row]
            joined = " | ".join(headers)
            has_code = any(h == "codigo" or h.startswith("codigo ") for h in headers)
            has_desc = any("descricao" in h for h in headers)
            has_unit = any(h in {"unidade", "un"} or "unidade" in h for h in headers)
            has_price = any(h in {"preco", "custo", "valor"} or "preco" in h or "custo" in h for h in headers)
            has_uf = any(h == "uf" for h in headers) or any(str(v or "").strip().upper() in UFS for v in row)
            if has_code and has_desc and has_unit and (has_price or has_uf):
                return idx, headers
        return -1, None

    def _kind_hint(self, sheet: str, workbook: str, headers: Sequence[str]) -> str | None:
        text = _norm(f"{sheet} {workbook} {' '.join(headers)}")
        if "compos" in text and "insumo" not in text: return "composition"
        if "insumo" in text and "compos" not in text: return "input"
        return None

    def _parse_row(self, row: Sequence[Any], headers: Sequence[str], kind_hint: str | None, release_id: str, workbook: str, sheet: str, uf_hint: str | None = None) -> list[SinapiItem]:
        values = list(row) + [None] * max(0, len(headers) - len(row))
        cols = {headers[i]: values[i] for i in range(len(headers))}
        code_col = next((h for h in headers if h == "codigo" or h.startswith("codigo ")), None)
        desc_col = next((h for h in headers if "descricao" in h), None)
        unit_col = next((h for h in headers if h in {"unidade","un"} or "unidade" in h), None)
        if not code_col or not desc_col or not unit_col:
            return []
        code = _code(cols.get(code_col)); desc = str(cols.get(desc_col) or "").strip(); unit = normalize_unit(cols.get(unit_col))
        if not code or not desc or not unit:
            return []
        kind = kind_hint
        type_col = next((h for h in headers if h in {"tipo","tipo de item","item"}), None)
        if type_col:
            t = _norm(cols.get(type_col))
            if "compos" in t: kind = "composition"
            elif "insumo" in t: kind = "input"
        if kind is None:
            ctext = _norm(code_col)
            if "compos" in ctext: kind = "composition"
            elif "insumo" in ctext: kind = "input"
        if kind is None:
            return []

        uf_col = next((h for h in headers if h == "uf"), None)
        price_col = next((h for h in headers if h in {"preco","custo","valor"} or "preco" in h or "custo" in h), None)
        out: list[SinapiItem] = []
        if uf_col and price_col:
            uf = str(cols.get(uf_col) or "").strip().upper(); price = _number(cols.get(price_col))
            if uf in UFS and price is not None:
                out.append(SinapiItem(release_id, kind, uf, code, desc, unit, price, workbook, sheet))
            return out
        if price_col and uf_hint:
            price = _number(cols.get(price_col))
            if price is not None:
                out.append(SinapiItem(release_id, kind, uf_hint, code, desc, unit, price, workbook, sheet))
            return out
        # Wide layout: one column per UF, including labels such as "Preço SP".
        for i, h in enumerate(headers):
            uf = h.upper() if h.upper() in UFS else _uf_hint(h)
            if uf in UFS:
                price = _number(values[i])
                if price is not None:
                    out.append(SinapiItem(release_id, kind, uf, code, desc, unit, price, workbook, sheet))
        return out


class SinapiService:
    def __init__(self, db: SinapiDatabase) -> None:
        self.db = db

    def list_releases(self) -> list[dict[str, Any]]:
        return [asdict(r) | {"composition_count": self.db.count_items(r.release_id, kind="composition"), "input_count": self.db.count_items(r.release_id, kind="input")} for r in self.db.list_releases()]

    def search_compositions(self, release_id: str, *, uf: str, query: str, limit: int = 20) -> list[dict[str, Any]]:
        tokens = [t for t in _norm(query).split() if len(t) >= 2]
        rows = self.db.items(release_id, kind="composition", uf=uf)
        def score(item: SinapiItem) -> tuple[float, str]:
            hay = _norm(f"{item.code} {item.description}")
            if not tokens: return (0.0, item.code)
            matched = sum(1 for token in tokens if token in hay)
            exact = 1 if _norm(query) and _norm(query) in hay else 0
            return (matched / len(tokens) + exact, item.code)
        ranked = [(score(item)[0], item) for item in rows]
        ranked = [pair for pair in ranked if not tokens or pair[0] > 0]
        ranked.sort(key=lambda x: (-x[0], x[1].description))
        return [asdict(item) | {"score": round(s, 4)} for s, item in ranked[:max(1, limit)]]

    def build_pricebook(self, release_id: str, *, uf: str, bdi_percent: float, mappings: Sequence[dict[str, Any]]) -> PriceBook:
        release = next((r for r in self.db.list_releases() if r.release_id == release_id), None)
        if release is None:
            raise KeyError("Publicação SINAPI não encontrada")
        comps: list[Composition] = []; cost_mappings: list[CostMapping] = []
        seen: set[str] = set()
        for raw in mappings:
            code = str(raw.get("composition_code") or "").strip()
            item = self.db.get_item(release_id, kind="composition", uf=uf, code=code)
            if item is None:
                raise ValueError(f"Composição SINAPI {code} não encontrada para {uf}")
            internal = f"SINAPI-{code}"
            if internal not in seen:
                comps.append(Composition(
                    code=internal, description=item.description, unit=item.unit,
                    components=(CostComponent(code=code, description=item.description, category="SINAPI", unit=item.unit, coefficient=1.0, unit_cost=item.price),),
                    waste_percent=float(raw.get("waste_percent") or 0.0),
                )); seen.add(internal)
            cost_mappings.append(CostMapping(
                composition_code=internal,
                ifc_class=str(raw.get("ifc_class") or "*"),
                quantity_kind=str(raw.get("quantity_kind") or ""),
                type_contains=raw.get("type_contains") or None,
                quantity_name_contains=raw.get("quantity_name_contains") or None,
                material_contains=raw.get("material_contains") or None,
                allow_geometry_fallback=bool(raw.get("allow_geometry_fallback", False)),
            ))
        provenance = {
            "release_id": release.release_id,
            "competence": release.competence,
            "uf": uf.upper(),
            "source_filename": release.source_filename,
            "source_sha256": release.source_sha256,
            "source_url": release.source_url,
            "source_provider": release.source_provider,
            "published_on": release.published_on,
            "fetched_at": release.fetched_at,
        }
        return PriceBook(
            "BRL", float(bdi_percent), "SINAPI", f"{release.competence}/{uf.upper()}",
            tuple(comps), tuple(cost_mappings), provenance,
        )

    def compare_releases(self, old_release_id: str, new_release_id: str, *, uf: str, codes: Sequence[str] | None = None) -> dict[str, Any]:
        old = {i.code: i for i in self.db.items(old_release_id, kind="composition", uf=uf)}
        new = {i.code: i for i in self.db.items(new_release_id, kind="composition", uf=uf)}
        selected = sorted(set(codes or (set(old) & set(new))))
        rows = []
        for code in selected:
            a, b = old.get(code), new.get(code)
            if not a or not b: continue
            delta = b.price - a.price
            pct = (delta / a.price * 100.0) if a.price else None
            rows.append({"code": code, "description": b.description, "unit": b.unit, "old_price": a.price, "new_price": b.price, "variation": delta, "variation_percent": pct})
        return {"old_release_id": old_release_id, "new_release_id": new_release_id, "uf": uf.upper(), "rows": rows}
